#!/usr/bin/env python3
"""Installed Qt window → mouse selection → conversion → KDE notification.

Native Qt verification is separate from Dolphin's actual context-menu test.
Run using the installed venv Python with -I in a graphical session.
"""
import hashlib
import json
import time
from pathlib import Path

from PySide6.QtCore import QPointF, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from convertix.cli import setup_logging
from convertix.detect import detect
import convertix.gui as gui


def main():
    project=Path(__file__).resolve().parents[1]
    examples=project/'examples/v2'
    if not examples.exists():
        raise SystemExit('Primero ejecuta scripts/make_v2_examples.py con el entorno del proyecto.')
    setup_logging();app=QApplication([]);app.setQuitOnLastWindowClosed(False)
    notifications=[];original_notify=gui.notify
    def observed(*args):
        result=original_notify(*args);notifications.append(result);return result
    gui.notify=observed
    cases=[(project/'examples/Documento de prueba.docx','pdf','docx-to-pdf'),
           (examples/'PDF digital con tabla.pdf','docx','pdf-to-docx'),
           (examples/'PDF escaneado.pdf','docx','ocr-to-docx'),
           (examples/'PDF escaneado.pdf','md','ocr-to-markdown'),
           (examples/'PDF digital con tabla.pdf','xlsx','pdf-to-xlsx'),
           (examples/'PDF mixto con portada vacía.pdf','pdf-ocr','searchable-pdf')]
    records=[]
    def walk(item):
        for child in item.childItems():
            yield child;yield from walk(child)
    for source,target,label in cases:
        before=hashlib.sha256(source.read_bytes()).hexdigest()
        controller=gui.Controller([source]);controller.prefs.notifications=True
        engine=gui.create_engine(controller);window=engine.rootObjects()[0];window.show()
        def wait(predicate,seconds=60):
            deadline=time.monotonic()+seconds
            while not predicate() and time.monotonic()<deadline:QTest.qWait(30)
            assert predicate(),controller.details
        assert QTest.qWaitForWindowExposed(window,5000)
        window.requestActivate();controller.scan();wait(lambda:controller.state!='scanning')
        assert controller.state=='ready',controller.details
        wait(lambda:not controller._jobs)
        index=next(i for i,o in enumerate(controller.options) if o['target']==target)
        window.setProperty('page',index//8);window.setProperty('selected',index%8);QTest.qWait(300)
        option=next(i for i in walk(window.contentItem()) if i.objectName()=='target_'+target)
        QTest.mouseClick(window,Qt.LeftButton,Qt.NoModifier,option.mapToScene(QPointF(option.width()/2,option.height()/2)).toPoint())
        wait(lambda:controller.state in ('done','error'));wait(lambda:not controller._jobs)
        assert controller.state=='done',controller.details
        output=Path(controller.records[0]['location']);assert detect(output).format==('pdf' if target=='pdf-ocr' else target)
        if target=='docx':
            from docx import Document
            assert any('Convertix' in p.text for p in Document(output).paragraphs)
        if target=='md':
            text=output.read_text();assert '# ' in text and 'Convertix' in text and '| ---' in text
        if target=='xlsx':
            from openpyxl import load_workbook
            book=load_workbook(output);assert book.active['B2'].value==42;book.close()
        assert hashlib.sha256(source.read_bytes()).hexdigest()==before
        window.grabWindow().save(str(project/f'docs/v2-completed-{label}.png'))
        records.append({'source':str(source),'target':target,'output':str(output),'original_sha256':before,
                        'original_unchanged':True,'mouse_selection':True,'notification_accepted':notifications[-1] is True})
        window.close();engine.deleteLater();QTest.qWait(100)
    result={'platform':app.platformName(),'cases':records,'dolphin_context_menu':'separate test; do not infer from this result'}
    (project/'docs/v2-desktop-smoke-result.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))


if __name__=='__main__':main()
