import hashlib
import io
import shutil
import threading
import zipfile
from pathlib import Path

import pytest
from PIL import Image, ImageDraw

from convertix.config import Preferences
from convertix.detect import detect
from convertix.models import ConversionError
from convertix.registry import routes
from convertix.service import convert

pymupdf = pytest.importorskip("pymupdf")
pytest.importorskip("docx")


@pytest.fixture(scope="session")
def pdf_samples(tmp_path_factory):
    folder = tmp_path_factory.mktemp("convertix-v2-pdf")
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((48, 55), "Informe Convertix", fontsize=24)
    page.insert_text((48, 95), "Documento en español: café, año y conversión.", fontsize=12)
    page.insert_text((48, 120), "English text: editable paragraphs and tables.", fontsize=12)
    page.insert_text((48, 150), "- Primer elemento de la lista", fontsize=12)
    page.insert_text((48, 175), "print('Convertix code')", fontname="cour", fontsize=11)
    page.insert_text((48, 200), "Documentación", fontsize=12)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(48,186,150,205), "uri": "https://example.org/docs"})
    xs=[48,245,380,540];ys=[240,280,320,360]
    for x in xs:page.draw_line((x,ys[0]),(x,ys[-1]),width=1)
    for y in ys:page.draw_line((xs[0],y),(xs[-1],y),width=1)
    rows=[["Nombre","Cantidad","Código"],["Café","42","0012"],["Té","7","=2+2"]]
    for r,row in enumerate(rows):
        for c,value in enumerate(row):page.insert_text((xs[c]+8,ys[r]+25),value,fontsize=12)
    image=Image.new('RGB',(120,80),'#99ccff');ImageDraw.Draw(image).ellipse((20,10,100,70),fill='#3355aa')
    buffer=io.BytesIO();image.save(buffer,'PNG')
    page.insert_image(pymupdf.Rect(380,100,500,180),stream=buffer.getvalue())
    page=document.new_page(width=595,height=842)
    page.insert_text((48,55),'Dos columnas',fontsize=24)
    for i in range(3):
        page.insert_text((48,110+i*80),f'Izquierda {i+1}: lectura ordenada.',fontsize=11)
        page.insert_text((320,110+i*80),f'Derecha {i+1}: segunda columna.',fontsize=11)
    digital=folder/'digital.pdf';document.save(digital)
    scan=pymupdf.open();page=scan.new_page(width=595,height=842)
    page.insert_image(page.rect,stream=document[0].get_pixmap(dpi=200).tobytes('png'))
    scanned=folder/'escaneado.pdf';scan.save(scanned)
    rotated=pymupdf.open();rotated.insert_pdf(scan);rotated[0].set_rotation(90);rotated.save(folder/'rotado.pdf')
    mixed=pymupdf.open();mixed.new_page();mixed.insert_pdf(document,from_page=0,to_page=0);mixed.insert_pdf(scan);mixed.save(folder/'mixto.pdf')
    empty=pymupdf.open();empty.new_page();empty.save(folder/'vacío.pdf')
    picture=pymupdf.open();p=picture.new_page();plain=Image.new('RGB',(80,80),'blue');b=io.BytesIO();plain.save(b,'PNG');p.insert_image(p.rect,stream=b.getvalue());picture.save(folder/'sin-texto.pdf')
    return {"digital":digital,"scanned":scanned,"mixed":folder/'mixto.pdf',"rotated":folder/'rotado.pdf',
            "blank":folder/'vacío.pdf',"no_text":folder/'sin-texto.pdf'}


def copy_sample(samples,kind,folder):
    path=folder/f"prueba {kind} café.pdf";shutil.copyfile(samples[kind],path);return path


@pytest.mark.integration
def test_digital_pdf_docx_editable(pdf_samples,tmp_path):
    from docx import Document
    source=copy_sample(pdf_samples,'digital',tmp_path);before=hashlib.sha256(source.read_bytes()).digest()
    result=convert(source,'docx');doc=Document(result.outputs[0])
    text='\n'.join(p.text for p in doc.paragraphs)
    assert 'Informe Convertix' in text and 'español' in text and 'English text' in text
    assert text.index('Izquierda 3')<text.index('Derecha 1')
    assert any(p.style.name.startswith('Heading') for p in doc.paragraphs)
    assert doc.tables and [[c.text for c in row.cells] for row in doc.tables[0].rows]==[
        ['Nombre','Cantidad','Código'],['Café','42','0012'],['Té','7','=2+2']]
    assert doc.inline_shapes and len(doc.sections)>=2
    assert any(s._sectPr.xpath('./w:cols')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num')=='2' for s in doc.sections if s._sectPr.xpath('./w:cols'))
    assert hashlib.sha256(source.read_bytes()).digest()==before


def test_pdf_page_classification(pdf_samples):
    assert detect(pdf_samples['digital']).pdf.kind=='digital'
    assert detect(pdf_samples['scanned']).pdf.kind=='scanned'
    mixed=detect(pdf_samples['mixed']).pdf
    assert mixed.kind=='mixed' and [p.needs_ocr for p in mixed.pages]==[False,False,True]
    assert detect(pdf_samples['blank']).pdf.kind=='no_text'
    assert detect(pdf_samples['no_text']).pdf.kind=='no_text'


@pytest.mark.integration
@pytest.mark.parametrize('kind',['scanned','mixed','rotated'])
def test_ocr_editable_docx_and_original(pdf_samples,tmp_path,kind):
    from docx import Document
    source=copy_sample(pdf_samples,kind,tmp_path);before=source.read_bytes()
    result=convert(source,'docx')
    doc=Document(result.outputs[0]);text=' '.join(p.text for p in doc.paragraphs)
    assert 'Convertix' in text and 'English' in text and 'español' in text
    assert len(text)>80
    # Only the digital page's separate illustration is retained, never a full-page scan.
    assert len(doc.inline_shapes)==(1 if kind=='mixed' else 0)
    assert source.read_bytes()==before


@pytest.mark.integration
def test_searchable_pdf_preserves_appearance_and_skips_existing_ocr(pdf_samples,tmp_path,monkeypatch):
    source=copy_sample(pdf_samples,'mixed',tmp_path);before=source.read_bytes()
    output=convert(source,'pdf-ocr').outputs[0]
    with pymupdf.open(source) as a,pymupdf.open(output) as b:
        assert len(a)==len(b)==3
        assert 'Convertix' in b[2].get_text() and not a[2].get_text().strip()
        assert a[1].get_text()==b[1].get_text()
        for pa,pb in zip(a,b):assert pa.get_pixmap(dpi=90).samples==pb.get_pixmap(dpi=90).samples
    profile=detect(output).pdf
    assert not profile.needs_ocr and any(p.kind=='ocr' for p in profile.pages)
    import convertix.adapters.ocr as ocr
    monkeypatch.setattr(ocr,'run',lambda *a,**k:pytest.fail('Existing OCR was processed twice'))
    ocr.searchable(detect(output),tmp_path/'unneeded.pdf',Preferences())
    assert source.read_bytes()==before


@pytest.mark.integration
@pytest.mark.parametrize('kind',['digital','scanned','mixed'])
def test_pdf_xlsx_real_tables(pdf_samples,tmp_path,kind):
    from openpyxl import load_workbook
    source=copy_sample(pdf_samples,kind,tmp_path)
    assert 'xlsx' in {r.target for r in routes(detect(source))}
    result=convert(source,'xlsx');book=load_workbook(result.outputs[0])
    assert len(book.worksheets)==(2 if kind=='mixed' else 1)
    values=list(book.active.values)
    assert len(values)==3 and len(values[0])==3
    assert values[1][1]==42 and str(values[1][2])=='0012'
    assert values[2][2]=='=2+2' and book.active['C3'].data_type=='s'
    book.close()


@pytest.mark.integration
@pytest.mark.parametrize('target',['md','html','odt','rtf'])
@pytest.mark.parametrize('kind',['digital','scanned','mixed'])
def test_pdf_other_editable_formats(pdf_samples,tmp_path,target,kind):
    source=copy_sample(pdf_samples,kind,tmp_path);before=source.read_bytes()
    result=convert(source,target)
    main=next(p for p in result.outputs if p.suffix=='.'+target)
    assert detect(main).format==target
    if target in ('md','html'):
        text=main.read_text();assert 'Convertix' in text and 'English' in text and 'Cantidad' in text
        if kind!='scanned':
            assert 'https://example.org/docs' in text and 'assets/' in text
        for asset in main.parent.glob('assets/*'):assert asset.stat().st_size>0
        assert len(list(main.parent.glob('assets/*')))==(0 if kind=='scanned' else 1)
        if target=='md':
            assert '# Informe' in text and '| ---' in text
            if kind!='scanned':assert '```' in text
        else:
            assert '<h1>' in text and '<table>' in text
            if kind!='scanned':assert '<code>' in text
    assert source.read_bytes()==before


def test_ocr_dependencies_are_explicit(pdf_samples,monkeypatch):
    import convertix.registry as registry
    import convertix.capabilities as capabilities
    monkeypatch.setattr(capabilities,'ocr_capability',lambda:capabilities.OcrCapability(None,()))
    source=detect(pdf_samples['scanned'])
    assert not {'docx','xlsx','md','html','odt','rtf','pdf-ocr','txt'} & {r.target for r in registry.routes(source)}
    assert not {'docx','pdf-ocr'} & {r.target for r in registry.routes(source,Preferences(ocr_languages='deu'))}
    mixed=detect(pdf_samples['mixed'])
    assert 'xlsx' not in {r.target for r in registry.routes(mixed)}


def test_no_table_no_spreadsheet(pdf_samples):
    doc=pymupdf.open();p=doc.new_page();p.insert_text((50,80),'Plain paragraph without a real table.');path=pdf_samples['digital'].parent/'paragraph.pdf';doc.save(path)
    assert 'xlsx' not in {r.target for r in routes(detect(path))}


def test_ocr_cancellation_cleans_everything(pdf_samples,tmp_path):
    from convertix.models import Cancelled
    source=copy_sample(pdf_samples,'scanned',tmp_path);before=source.read_bytes();cancel=threading.Event();events=[]
    def progress(message):
        events.append(message)
        if message.startswith('OCR'):cancel.set()
    with pytest.raises(Cancelled):convert(source,'docx',cancel=cancel,progress=progress)
    assert any(m.startswith('OCR') for m in events)
    assert source.read_bytes()==before and not list(tmp_path.glob('.convertix-*')) and not list(tmp_path.glob('*.docx'))


@pytest.mark.parametrize('language,phrase',[('spa','español'),('eng','English')])
def test_ocr_language_preference(pdf_samples,tmp_path,language,phrase):
    source=copy_sample(pdf_samples,'scanned',tmp_path)
    text=convert(source,'txt',Preferences(ocr_languages=language)).outputs[0].read_text()
    assert phrase in text


def test_two_independent_tables_become_two_sheets(tmp_path):
    from openpyxl import load_workbook
    doc=pymupdf.open();page=doc.new_page()
    for offset,label in [(100,'Primera'),(400,'Segunda')]:
        for x in [50,220,390]:page.draw_line((x,offset),(x,offset+100))
        for y in [offset,offset+50,offset+100]:page.draw_line((50,y),(390,y))
        for r,row in enumerate([[label,'Valor'],['Dato','17']]):
            for c,value in enumerate(row):page.insert_text((60+c*170,offset+30+r*50),value)
    source=tmp_path/'tablas independientes.pdf';doc.save(source)
    result=convert(source,'xlsx');book=load_workbook(result.outputs[0])
    assert len(book.worksheets)==2
    assert book.worksheets[0]['A1'].value=='Primera' and book.worksheets[1]['A1'].value=='Segunda'
    book.close()


def test_same_page_digital_header_and_scanned_body(pdf_samples,tmp_path):
    source=tmp_path/'página mixta.pdf';doc=pymupdf.open();page=doc.new_page()
    page.insert_text((40,50),'Cabecera digital que debe permanecer sin duplicar ni volver a reconocer.',fontsize=10)
    with pymupdf.open(pdf_samples['scanned']) as scanned:
        page.insert_image(pymupdf.Rect(40,90,555,820),stream=scanned[0].get_pixmap(dpi=200).tobytes('png'))
    doc.save(source)
    profile=detect(source).pdf
    assert profile.kind=='mixed' and profile.pages[0].needs_ocr
    text=convert(source,'txt').outputs[0].read_text()
    assert text.count('Cabecera digital')==1 and 'English' in text


def test_cancellable_probe_profiles_pdf(pdf_samples):
    from convertix.probe import probe
    from convertix.models import Cancelled
    assert probe(pdf_samples['scanned']).pdf.kind=='scanned'
    event=threading.Event();event.set()
    with pytest.raises(Cancelled):probe(pdf_samples['digital'],cancel=event)


def test_previous_ocr_tables_need_no_new_recognition(pdf_samples,tmp_path,monkeypatch):
    from openpyxl import load_workbook
    import convertix.capabilities as capabilities
    from convertix.adapters.pdf_documents import convert as adapter
    source=copy_sample(pdf_samples,'scanned',tmp_path)
    searchable=convert(source,'pdf-ocr').outputs[0]
    inspected=detect(searchable)
    assert not inspected.pdf.needs_ocr and inspected.pdf.scan_tables==1
    monkeypatch.setattr(capabilities,'ocr_capability',lambda:capabilities.OcrCapability(None,()))
    assert 'xlsx' in {r.target for r in routes(inspected)}
    folder=tmp_path/'sin motor OCR';folder.mkdir()
    adapter(inspected,'xlsx',folder,Preferences())
    book=load_workbook(folder/'document.xlsx');assert book.active['B2'].value==42;book.close()
