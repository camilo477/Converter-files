"""Editable DOCX serialization. Other local serializers share this document model."""
import io
import html
import re
from pathlib import Path

from .models import ConversionError
from .pdf_layout import LayoutPage, list_style


def resources(pages: list[LayoutPage], output: Path) -> dict[int, str]:
    assets = {}
    for page_number, page in enumerate(pages, 1):
        image_number = 0
        for element in page.elements:
            if element.kind == "image":
                image_number += 1
                folder = output/"assets"; folder.mkdir(exist_ok=True)
                name = f"page-{page_number:03}-image-{image_number:03}.png"
                (folder/name).write_bytes(element.image)
                assets[id(element)] = "assets/"+name
    return assets


def md_escape(value: str) -> str:
    return re.sub(r"([\\`*_\[\]<>])", r"\\\1", value)


def markdown(pages: list[LayoutPage], output: Path) -> None:
    images = resources(pages, output)
    parts = []
    for number, page in enumerate(pages, 1):
        if number > 1:
            parts.append(f"<!-- Página {number} -->")
        for e in page.elements:
            if e.kind == "table":
                rows = [[v.replace("|", r"\|").replace("\n", " ") for v in row] for row in e.rows]
                columns = max(map(len, rows))
                rows = [row+[""]*(columns-len(row)) for row in rows]
                parts.append("\n".join(["| "+" | ".join(rows[0])+" |", "| "+" | ".join(["---"]*columns)+" |"]+
                                       ["| "+" | ".join(row)+" |" for row in rows[1:]]))
            elif e.kind == "image":
                parts.append(f"![Imagen de la página {number}]({images[id(e)]})")
            elif any(s.mono for s in e.spans):
                fence = "`"*max(3, max((len(m[0])+1 for m in re.finditer(r"`+", e.text)), default=3))
                parts.append(f"{fence}\n{e.text}\n{fence}")
            else:
                text = ""
                for span in e.spans:
                    value = md_escape(span.text)
                    if span.bold and value.strip(): value = "**"+value+"**"
                    if span.italic and value.strip(): value = "*"+value+"*"
                    if span.url:
                        from urllib.parse import quote
                        value = f"[{value}]({quote(span.url, safe=':/?#&=%@+;,-._~')})"
                    text += value
                if e.level:
                    text = "#"*e.level+" "+text
                elif list_style(e.text):
                    text = re.sub(r"^[•●◦]\s*", "- ", text)
                parts.append(text)
    result = "\n\n".join(parts).strip()+"\n"
    if not result.strip():
        raise ConversionError("No se obtuvo contenido para Markdown.")
    (output/"document.md").write_text(result)


def html_document(pages: list[LayoutPage], output: Path) -> None:
    images = resources(pages, output)
    parts = ['<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Convertix</title>',
             '<style>body{font:12pt sans-serif;max-width:70em;margin:2em auto;padding:1em}section{break-after:page}table{border-collapse:collapse}td,th{border:1px solid #888;padding:.5em}img{max-width:100%;height:auto}pre{white-space:pre-wrap}</style></head><body>']
    for number, page in enumerate(pages, 1):
        parts.append(f'<section aria-label="Página {number}">')
        for e in page.elements:
            if e.kind == "table":
                parts.append('<table>')
                for index, row in enumerate(e.rows):
                    tag = "th" if index == 0 else "td"
                    parts.append('<tr>'+''.join(f'<{tag}>{html.escape(v)}</{tag}>' for v in row)+'</tr>')
                parts.append('</table>')
            elif e.kind == "image":
                parts.append(f'<img src="{images[id(e)]}" alt="Imagen de la página {number}">')
            elif any(s.mono for s in e.spans):
                parts.append('<pre><code>'+html.escape(e.text)+'</code></pre>')
            else:
                content = ""
                for span in e.spans:
                    value = html.escape(span.text)
                    if span.bold: value = '<strong>'+value+'</strong>'
                    if span.italic: value = '<em>'+value+'</em>'
                    if span.url: value = '<a href="'+html.escape(span.url, quote=True)+'">'+value+'</a>'
                    content += value
                tag = f'h{e.level}' if e.level else 'p'
                parts.append(f'<{tag}>{content}</{tag}>')
        parts.append('</section>')
    parts.append('</body></html>')
    (output/'document.html').write_text('\n'.join(parts))


def to_xlsx(pages: list[LayoutPage], destination: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment
    book = Workbook(); book.remove(book.active)
    for p, page in enumerate(pages, 1):
        number = 0
        for element in page.elements:
            if element.kind != "table":
                continue
            number += 1
            sheet = book.create_sheet(f'P{p}_Tabla{number}')
            for r, row in enumerate(element.rows, 1):
                for c, text in enumerate(row, 1):
                    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
                    cell = sheet.cell(r, c)
                    # Preserve ambiguous identifiers and never activate PDF text as an Excel formula.
                    if re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?', text.strip()) and len(text.strip()) < 16:
                        cell.value = float(text) if '.' in text else int(text)
                    else:
                        cell.value = text
                        cell.data_type = 's'
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    if r == 1: cell.font = Font(bold=True)
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or '')) for c in column)+2))
    if not book.worksheets:
        raise ConversionError('No se detectaron tablas con filas, columnas y datos utilizables. No se creó una hoja vacía.')
    book.save(destination)
    check = load_workbook(destination, read_only=True, data_only=False)
    try:
        if not any(cell.value is not None for sheet in check for row in sheet for cell in row):
            raise ConversionError('El XLSX no contiene datos extraídos.')
    finally:
        check.close()


def to_docx(pages: list[LayoutPage], destination: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    document = Document()
    document.core_properties.title = "Documento convertido con Convertix"
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    text_count = 0
    for index, page in enumerate(pages):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width = Pt(page.width)
        section.page_height = Pt(page.height)
        section.left_margin = section.right_margin = Pt(32)
        section.top_margin = section.bottom_margin = Pt(32)
        current_columns = 1
        previous_column = 0
        for element in page.elements:
            columns = 2 if page.columns == 2 and element.column else 1
            if columns != current_columns:
                section = document.add_section(WD_SECTION_START.CONTINUOUS)
                cols = section._sectPr.find(qn("w:cols"))
                if cols is None:
                    cols = OxmlElement("w:cols"); section._sectPr.append(cols)
                cols.set(qn("w:num"), str(columns))
                cols.set(qn("w:space"), "360")
                current_columns = columns
            if columns == 2 and previous_column == -1 and element.column == 1:
                document.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
            previous_column = element.column
            if element.kind == "table":
                table = document.add_table(rows=len(element.rows), cols=max(map(len, element.rows)))
                table.style = "Table Grid"
                for r, values in enumerate(element.rows):
                    for c, value in enumerate(values):
                        table.cell(r, c).text = value
                        text_count += len(value.strip())
            elif element.kind == "image":
                available = (page.width-64)/(2 if columns == 2 else 1)
                width = max(12, min(element.bbox[2]-element.bbox[0], available))
                document.add_picture(io.BytesIO(element.image), width=Pt(width))
            else:
                style = f"Heading {element.level}" if element.level else list_style(element.text) or "Normal"
                paragraph = document.add_paragraph(style=style)
                if list_style(element.text):
                    # The literal marker already carries the numbering; don't invent a new sequence.
                    paragraph.style = document.styles["Normal"]
                for span in element.spans:
                    run = paragraph.add_run(span.text)
                    if span.url:
                        from docx.opc.constants import RELATIONSHIP_TYPE
                        link = OxmlElement("w:hyperlink")
                        link.set(qn("r:id"), paragraph.part.relate_to(span.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                        link.append(run._r)
                        paragraph._p.append(link)
                        run.underline = True
                    run.bold = span.bold
                    run.italic = span.italic
                    run.font.name = "Consolas" if span.mono else {"Helvetica": "Arial", "Times-Roman": "Times New Roman"}.get(span.font, span.font)
                    run.font.size = Pt(max(5, min(72, span.size)))
                    run.font.color.rgb = RGBColor.from_string(f"{span.color & 0xffffff:06x}")
                    text_count += len(span.text.strip())
    if not text_count:
        raise ConversionError("No se obtuvo texto editable. El PDF puede necesitar OCR o no contener texto.")
    document.save(destination)
    check = Document(destination)
    actual = "".join(p.text for p in check.paragraphs) + "".join(c.text for t in check.tables for r in t.rows for c in r.cells)
    if not actual.strip():
        raise ConversionError("El DOCX generado no contiene texto editable.")
